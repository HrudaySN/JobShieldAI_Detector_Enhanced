import json
from datetime import datetime
from flask import Blueprint, jsonify, render_template, request
from flask_login import current_user, login_required

from extensions import db
from job_analyzer import model
from resume_analyzer import model as resume_model
from file_processor import extract_text as extract_resume_text
from models import JobAnalysis, ResumeScreening

main_bp = Blueprint("main", __name__)


@main_bp.route("/")
def home():
    return render_template("home.html", active_page="home")


@main_bp.route("/detector")
@login_required
def detector():
    return render_template("detector.html", active_page="detector", training_examples=model.examples)


@main_bp.route("/api/extract-job-file", methods=["POST"])
@login_required
def extract_job_file():
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"error": "Please select a job file."}), 400

    filename = file.filename.lower()
    allowed = (".txt", ".pdf", ".docx", ".png", ".jpg", ".jpeg", ".webp")
    if not filename.endswith(allowed):
        return jsonify({"error": "Supported files: TXT, PDF, DOCX, PNG, JPG, JPEG or WEBP."}), 400

    raw = file.read()
    if len(raw) > 5 * 1024 * 1024:
        return jsonify({"error": "File is larger than the 5MB limit."}), 400

    try:
        if filename.endswith(".txt"):
            text = raw.decode("utf-8", errors="ignore")
        elif filename.endswith(".pdf"):
            from PyPDF2 import PdfReader
            import io
            reader = PdfReader(io.BytesIO(raw))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif filename.endswith(".docx"):
            from docx import Document
            import io
            doc = Document(io.BytesIO(raw))
            text = "\n".join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(cell.text for cell in row.cells)
        else:
            try:
                from PIL import Image
                import io
                import pytesseract
                image = Image.open(io.BytesIO(raw))
                text = pytesseract.image_to_string(image)
            except Exception as exc:
                return jsonify({"error": "Image text extraction is unavailable. Install Tesseract OCR and try again, or paste the job description."}), 422

        text = "\n".join(line.strip() for line in text.splitlines() if line.strip()).strip()
        if len(text.split()) < 4:
            return jsonify({"error": "The selected file did not contain enough readable job text. Try a clearer file or paste the description."}), 422
        return jsonify({"text": text[:5000], "filename": file.filename})
    except Exception as exc:
        return jsonify({"error": f"Could not read the selected file: {exc}"}), 422


@main_bp.route("/api/analyze-job", methods=["POST"])
@login_required
def analyze_job():
    payload = request.get_json(silent=True) or {}
    try:
        return jsonify(model.analyze(payload.get("text", ""), payload))
    except ValueError as error:
        return jsonify({"error": str(error)}), 400


@main_bp.route("/api/save-analysis", methods=["POST"])
@login_required
def save_analysis():
    payload = request.get_json(silent=True) or {}
    result = payload.get("result") or {}
    text = (payload.get("text") or "").strip()
    if not text or not result:
        return jsonify({"error": "Nothing to save yet."}), 400
    record = JobAnalysis(
        user_id=current_user.id,
        job_text=text[:10000],
        job_url=(payload.get("job_url") or "")[:1000] or None,
        recruiter_email=(payload.get("recruiter_email") or "")[:255] or None,
        company_website=(payload.get("company_website") or "")[:1000] or None,
        risk_score=int(result.get("risk_score", 0)),
        risk_level=result.get("risk_level", "Medium"),
        verdict=result.get("verdict", "Needs review"),
        result_json=json.dumps(result),
    )
    db.session.add(record)
    db.session.commit()
    return jsonify({"ok": True, "id": record.id})


@main_bp.route("/api/analysis/<int:analysis_id>")
@login_required
def get_analysis(analysis_id):
    record = JobAnalysis.query.filter_by(id=analysis_id, user_id=current_user.id).first_or_404()
    return jsonify({
        "id": record.id,
        "text": record.job_text,
        "job_url": record.job_url,
        "recruiter_email": record.recruiter_email,
        "company_website": record.company_website,
        "created_at": record.created_at.isoformat(),
        "result": json.loads(record.result_json),
    })


@main_bp.route("/resume-screening")
@login_required
def resume_screen():
    return render_template("resume.html", active_page="resume")

@main_bp.route("/api/extract-resume-file", methods=["POST"])
@login_required
def extract_resume_file():
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"error": "Please select a resume file."}), 400

    raw_len = len(file.read())
    file.seek(0)
    if raw_len > 5 * 1024 * 1024:
        return jsonify({"error": "File is larger than the 5MB limit."}), 400

    try:
        text = extract_resume_text(file)
        return jsonify({"text": text[:8000], "filename": file.filename})
    except ValueError as error:
        return jsonify({"error": str(error)}), 422
    except Exception:
        return jsonify({"error": "Could not read that file. Please try a PDF or DOCX resume."}), 422


@main_bp.route("/api/screen-resume", methods=["POST"])
@login_required
def screen_resume():
    payload = request.get_json(silent=True) or {}
    try:
        result = resume_model.analyze(payload.get("text", ""), payload.get("target_role"))
        return jsonify(result)
    except ValueError as error:
        return jsonify({"error": str(error)}), 400


@main_bp.route("/api/save-resume-screening", methods=["POST"])
@login_required
def save_resume_screening():
    payload = request.get_json(silent=True) or {}
    result = payload.get("result") or {}
    text = (payload.get("text") or "").strip()
    if not text or not result:
        return jsonify({"error": "Nothing to save yet."}), 400
    record = ResumeScreening(
        user_id=current_user.id,
        resume_text=text[:10000],
        filename=(payload.get("filename") or "")[:255] or None,
        target_role=(result.get("target_role") or "")[:200] or None,
        fit_score=int(result.get("fit_score", 0)),
        skills_detected=json.dumps(result.get("skills_detected", [])),
        skills_missing=json.dumps(result.get("skills_missing", [])),
        result_json=json.dumps(result),
    )
    db.session.add(record)
    db.session.commit()
    return jsonify({"ok": True, "id": record.id})


@main_bp.route("/api/resume-screening/<int:screening_id>")
@login_required
def get_resume_screening(screening_id):
    record = ResumeScreening.query.filter_by(id=screening_id, user_id=current_user.id).first_or_404()
    return jsonify({
        "id": record.id,
        "text": record.resume_text,
        "filename": record.filename,
        "created_at": record.created_at.isoformat(),
        "result": json.loads(record.result_json),
    })


@main_bp.route("/dashboard")
@login_required
def dashboard():
    analyses = JobAnalysis.query.filter_by(user_id=current_user.id)
    total = analyses.count()
    high = analyses.filter(JobAnalysis.risk_level == "High").count()
    medium = analyses.filter(JobAnalysis.risk_level == "Medium").count()
    low = analyses.filter(JobAnalysis.risk_level == "Low").count()
    recent = analyses.order_by(JobAnalysis.created_at.desc()).limit(5).all()
    return render_template("dashboard.html", active_page="dashboard", stats={"total": total, "high": high, "medium": medium, "low": low}, recent=recent)


@main_bp.route("/history")
@login_required
def history():
    analyses = JobAnalysis.query.filter_by(user_id=current_user.id).order_by(JobAnalysis.created_at.desc()).all()
    return render_template("history.html", active_page="history", analyses=analyses)


@main_bp.route("/support")
def support():
    return render_template("support.html", active_page="support")
